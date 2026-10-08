"""Разбор документов базы знаний на фрагменты с сохранением номеров статей, частей и пунктов.

Три вида источников:
  * закон или кодекс (data/raw/law/*.txt, structure: law) — строка «Статья 4. …» открывает статью,
    «2. …» — часть (в НК РФ — пункт), «8) …» — пункт части (в НК РФ — подпункт). Ссылка в ответе
    собирается из этих номеров: «422-ФЗ, ст. 4, ч. 2, п. 8»;
  * разъяснения ФНС (structure: faq) — «Вопрос: …», ответ и «Источник: …» — статья закона,
    на которую ссылается ФНС; ссылка в ответе — вопрос ФНС;
  * PDF и DOCX (письма ФНС, памятки) — как в исходном прототипе: строка с номером пункта
    («4.2», «12.») открывает пункт, короткая строка с номером раздела — заголовок.

Пункты собираются во фрагменты удобного для поиска размера: мелкие соседние пункты одной статьи
объединяются, длинные режутся с перекрытием, но номера всегда остаются при фрагменте — они и есть
ссылка в ответе. Пункт из перечня («8) налогоплательщики, у которых доходы… превысили 2,4 миллиона»)
получает начало своей части («Не вправе применять специальный налоговый режим: …») — иначе
из фрагмента непонятно, что это запрет.
"""
import re
from collections import Counter
from dataclasses import asdict, dataclass, field
from pathlib import Path

import pymupdf
import yaml
from docx import Document as DocxDocument

from .config import settings
from .text import join_hyphenated, normalize

CLAUSE_RE = re.compile(r"^(?:п\.\s*)?([1-9]\d{0,2}(?:\.\d{1,3}){0,4})(\.?)\s+(?=[А-ЯЁA-Z«\"(])")
ROMAN_RE = re.compile(r"^([IVXL]{1,6})\.\s+[А-ЯЁA-Z]")
ANNEX_RE = re.compile(r"^(Приложение\s+(?:N|№)?\s*[\dА-Я]+)", re.IGNORECASE)

TARGET, MAX_LEN, OVERLAP = 900, 1400, 150
LAW_TARGET = 700          # статьи закона режутся мельче: ссылка на часть и пункт точнее, чем на всю статью


@dataclass
class Chunk:
    id: str
    doc_id: str
    doc_title: str
    doc_kind: str
    access: str
    clause: str | None           # часть статьи (пункт в НК РФ) или пункт документа; диапазон — «8–11»
    heading: str | None          # название статьи, вопрос ФНС или раздел документа
    page_start: int | None
    page_end: int | None
    text: str
    article: str | None = None   # номер статьи закона: «4», «346.13»
    item: str | None = None      # пункт части (подпункт в НК РФ): «8», «1–7»
    short: str | None = None     # краткое название документа для ссылки: «422-ФЗ», «НК РФ»
    url: str | None = None       # где прочитать статью или разъяснение в оригинале
    topic: str | None = None     # тема документа или статьи для поиска: у 456-ФЗ в тексте нет слов «самозанятый» и
    #                              «больничный», у статьи 1 422-ФЗ из заголовка не видно, о каком эксперименте речь
    part_word: str = "ч."        # как называется часть статьи в ссылке («ч.» или «п.»)
    item_word: str = "п."        # и пункт части («п.» или «пп.»)

    @property
    def search_text(self) -> str:
        """Текст для индексации: название документа, номер и название статьи помогают поиску."""
        if self.article:
            head = f"{self.short or self.doc_title}. Статья {self.article}. {self.heading or ''}."
        else:
            head = f"{self.short or self.doc_title}. {self.heading or ''}."
        if self.topic:
            head = f"{head} ({self.topic})"
        return f"{head}\n{self.text}"

    @property
    def ref(self) -> str:
        """«ст. 4, ч. 2, п. 8» — место в документе, без названия документа."""
        parts = []
        if self.article:
            parts.append(f"ст. {self.article}")
        if self.clause:
            parts.append(f"{self.part_word if self.article else 'п.'} {self.clause}")
        if self.item:
            parts.append(f"{self.item_word} {self.item}")
        return ", ".join(parts)

    def to_dict(self) -> dict:
        return asdict(self)


@dataclass
class _Unit:
    clause: str | None
    heading: str | None
    page_start: int | None
    page_end: int | None
    lines: list
    pages: list = field(default_factory=list)   # страница каждой строки

    @property
    def text(self) -> str:
        return " ".join(self.lines).strip()

    def spans(self) -> list[tuple[int, int | None]]:
        """Начало каждой строки в text и её страница."""
        out, pos = [], 0
        for line, page in zip(self.lines, self.pages):
            out.append((pos, page))
            pos += len(line) + 1
        return out


# ---------- чтение файлов ----------

def ocr_available() -> bool:
    return (settings.tessdata_dir / "rus.traineddata").exists()


def _page_text(page, stats: dict) -> str:
    """Текстовый слой страницы, а для скана — распознанный текст (Tesseract встроен в PyMuPDF,
    языки — tools/tessdata, скачиваются scripts/setup.py --ocr)."""
    # естественный порядок блоков: при sort=True колонки двухколоночной вёрстки перемешиваются
    text = page.get_text("text")
    if len(text.strip()) >= 20 or not page.get_images():
        return text
    if not ocr_available():
        stats["scanned_pages"] = stats.get("scanned_pages", 0) + 1
        return text
    tp = page.get_textpage_ocr(language="rus+eng", dpi=300, full=True, tessdata=str(settings.tessdata_dir))
    stats["ocr_pages"] = stats.get("ocr_pages", 0) + 1
    return page.get_text("text", textpage=tp)


def _pdf_pages(path: Path, stats: dict | None = None) -> list[tuple[int, list[str]]]:
    stats = {} if stats is None else stats
    doc = pymupdf.open(path)
    pages = []
    for i, page in enumerate(doc):
        raw = normalize(_page_text(page, stats))
        lines = [l.strip() for l in raw.split("\n")]
        pages.append((i + 1, [l for l in lines if l]))
    # повторяющиеся колонтитулы и номера страниц
    if len(pages) >= 4:
        freq = Counter(l for _, lines in pages for l in set(lines) if len(l) < 80)
        boiler = {l for l, c in freq.items() if c >= max(3, 0.3 * len(pages))}
    else:
        boiler = set()
    cleaned = []
    for no, lines in pages:
        # номер страницы — и «Страница 1 из 2» в коротком документе (иначе он прилипает к последнему пункту)
        lines = [l for l in lines if l not in boiler and not re.fullmatch(
            r"\d{1,4}|-\s*\d{1,4}\s*-|(?:Страница|Стр\.|Лист)\s*\d{1,4}(?:\s*(?:из|/)\s*\d{1,4})?", l, re.I)]
        cleaned.append((no, join_hyphenated(lines)))
    return cleaned


def _docx_pages(path: Path) -> list[tuple[int | None, list[str]]]:
    doc = DocxDocument(path)
    lines = []
    for p in doc.paragraphs:
        text = normalize(p.text).strip()
        if text:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [normalize(c.text).strip() for c in row.cells]
            row_text = " | ".join(dict.fromkeys(c for c in cells if c))
            if row_text:
                lines.append(row_text)
    return [(None, lines)]


# ---------- структура: пункты и заголовки ----------

def _is_heading(line: str, number: str | None, dotted: bool) -> bool:
    if ROMAN_RE.match(line) or ANNEX_RE.match(line):
        return True
    if number and "." not in number:
        # «4 Порядок проведения» — заголовок раздела ГОСТа: номер без точки и короткая строка.
        # «99. Перед хранением…» — пункт приказа, даже если строка оборвана переносом.
        return not dotted and len(line) < 90 and not line.rstrip().endswith((".", ";", ":"))
    letters = [c for c in line if c.isalpha()]
    return 8 < len(line) < 90 and len(letters) > 5 and all(c.isupper() for c in letters)


def _split_inline_clauses(lines: list[str], last_top):
    """В PDF приказов пункт может начинаться в середине строки:
    «…ремонт инструмента. 117. Работать с …». Строка разрезается, если внутри неё
    стоит номер следующего по порядку пункта. Генератор: номер последнего пункта
    обновляется по мере разбора, поэтому проверка идёт по актуальному значению."""
    for line in lines:
        while True:
            n = last_top() + 1
            m = re.search(rf"(?<=[.;:)»\"])\s+(?:{n}|{n + 1}|{n + 2})\.\s+(?=[А-ЯЁ])", line)
            if not m or m.start() < 3:
                yield line
                break
            yield line[:m.start()].strip()
            line = line[m.start():].strip()


def _units(pages) -> list[_Unit]:
    units: list[_Unit] = []
    heading = None
    last_top = 0
    current = _Unit(None, None, None, None, [])

    def flush():
        nonlocal current
        if current.lines:
            units.append(current)
        current = _Unit(current.clause, heading, None, None, [])

    for page_no, lines in pages:
        for line in _split_inline_clauses(lines, lambda: last_top):
            m = CLAUSE_RE.match(line)
            number = m.group(1) if m else None
            dotted = bool(m and m.group(2))
            if number and "." not in number:
                # одиночные номера принимаем только по порядку (допуская пропуск
                # до двух нераспознанных), чтобы «220 В» в тексте не стал пунктом 220
                value = int(number)
                if not (value == 1 or last_top < value <= last_top + 3):
                    number = None
            if _is_heading(line, number, dotted):
                flush()
                heading = line[:120]
                if number and "." not in number:
                    last_top = int(number)
                current = _Unit(number, heading, page_no, page_no, [])
                continue
            if number:
                flush()
                if "." not in number:
                    last_top = int(number)
                current = _Unit(number, heading, page_no, page_no, [line], [page_no])
                continue
            if current.page_start is None:
                current.page_start = page_no
                current.heading = heading
            current.page_end = page_no
            current.lines.append(line)
            current.pages.append(page_no)
    flush()
    return [u for u in units if u.text]


def _split_long(text: str) -> list[tuple[int, str]]:
    """Куски длинного текста с перекрытием и начало каждого куска в тексте."""
    parts, start = [], 0
    while start < len(text):
        end = min(start + TARGET, len(text))
        if end < len(text):
            cut = text.rfind(". ", start + TARGET // 2, end)
            end = cut + 1 if cut > 0 else end
        raw = text[start:end]
        parts.append((start + len(raw) - len(raw.lstrip()), raw.strip()))
        if end >= len(text):
            break
        # перекрытие начинается с начала предложения (или хотя бы слова), а не с середины слова:
        # иначе фрагмент начинается с «ники из числа…» и поиск его не узнаёт
        back = max(end - OVERLAP, start + 1)
        sentence = text.find(". ", back, end)
        if sentence != -1:
            start = sentence + 2
        else:
            space = text.find(" ", back, end)
            start = space + 1 if space != -1 else back
    return parts


def _chunks(units: list[_Unit], meta: dict) -> list[Chunk]:
    chunks: list[Chunk] = []
    buf: list[_Unit] = []

    def emit(group: list[_Unit]):
        text = "\n".join(u.text for u in group)
        clauses = [u.clause for u in group if u.clause]
        if not clauses:
            clause = None
        elif clauses[0] == clauses[-1]:
            clause = clauses[0]
        else:
            clause = f"{clauses[0]}–{clauses[-1]}"
        pages = [p for u in group for p in (u.page_start, u.page_end) if p]
        base = dict(doc_id=meta["id"], doc_title=meta["title"], doc_kind=meta["kind"], access=meta["access"],
                    clause=clause, heading=group[0].heading,
                    page_start=min(pages) if pages else None, page_end=max(pages) if pages else None)
        # страница каждой строки по её началу в общем тексте: у куска длинного пункта —
        # свои страницы, а не весь пункт («стр. 125», а не «стр. 124–126»)
        line_pages, offset = [], 0
        for u in group:
            line_pages += [(offset + pos, page) for pos, page in u.spans() if page]
            offset += len(u.text) + 1
        pieces = _split_long(text) if len(text) > MAX_LEN else [(0, text)]
        opening = re.sub(r"^\s*\d{1,3}(?:\.\d{1,3})*\.?\s+", "", pieces[0][1])     # без номера пункта
        lead = re.split(r"(?<=[.;:])\s", opening, maxsplit=1)[0][:200]
        for k, (found, piece) in enumerate(pieces):
            # продолжение длинного пункта получает в заголовок его начало: иначе фрагмент «из числа
            # административно-технического персонала…» не говорит, о чём он, и поиск его не узнаёт
            heading = base["heading"] if k == 0 else f"{base['heading'] or ''} — продолжение: {lead}".strip(" —")
            own = {**base, "heading": heading}
            if len(pieces) > 1:
                # строка, в которой кусок начинается, и все строки, начатые внутри него
                first = max((pos for pos, _ in line_pages if pos <= found), default=0)
                span = [p for pos, p in line_pages if first <= pos < found + len(piece)]
                if span:
                    own.update(page_start=min(span), page_end=max(span))
            chunks.append(Chunk(id=f"{meta['id']}#{len(chunks)}", text=piece, **own))

    # Пункты одного раздела объединяются, пока фрагмент не превысит целевой размер:
    # короткий раздел («2 Классификация отказов» из четырёх пунктов) остаётся целым.
    for unit in units:
        same_section = buf and buf[-1].heading == unit.heading
        size = sum(len(u.text) for u in buf)
        if buf and (not same_section or size + len(unit.text) > TARGET):
            emit(buf)
            buf = []
        buf.append(unit)
    if buf:
        emit(buf)
    return chunks


# ---------- каталог документов ----------

ARCHIVE = "archive"   # группа доступа отменённых редакций: её нет ни у одной роли
DRAFT = "admin"       # проект новой редакции до утверждения видит только администратор


def is_cancelled(info: dict) -> bool:
    """Статус в каталоге: «действует» (по умолчанию) или «отменён» (с полем replaced_by)."""
    return str(info.get("status", "действует")).lower().startswith("отмен")


def load_catalog(raw_dir: Path) -> dict[str, dict]:
    """Названия, группы доступа и статусы (действует / отменён) из data/sources.yaml и data/raw/*/catalog.yaml."""
    catalog = {}
    for path in [raw_dir.parent / "sources.yaml", *raw_dir.glob("*/catalog.yaml")]:
        if path.exists():
            for d in yaml.safe_load(path.read_text(encoding="utf-8"))["documents"]:
                catalog[d["id"]] = d
    # новая редакция, загруженная с пометкой «заменяет X», отменяет X — только после утверждения:
    # черновик с ошибкой не должен подменять действующую редакцию
    for d in list(catalog.values()):
        old = catalog.get(d.get("replaces"))
        if old is not None and not is_cancelled(d) and not is_draft_edition(d):
            old.update(status="отменён", replaced_by=d["id"])
    return catalog


def is_draft_edition(info: dict) -> bool:
    """Загруженная новая редакция, которую ещё не утвердили (у документов из sources.yaml поля approved нет)."""
    return bool(info.get("replaces")) and not info.get("approved", True)


def access_of(info: dict) -> str:
    if is_cancelled(info):
        return ARCHIVE
    return DRAFT if is_draft_edition(info) else info.get("access", "all")


# ---------- законы и разъяснения ФНС (текстовые файлы из scripts/fetch_laws.py) ----------

ARTICLE_RE = re.compile(r"^Статья\s+(\d{1,3}(?:\.\d{1,2}){0,2})\.\s*(.*)$")
PART_RE = re.compile(r"^(\d{1,3}(?:\.\d{1,2})?)\.\s+(.+)$")         # «2. …», «1.2. …», «2.1. …»
ITEM_RE = re.compile(r"^(\d{1,3}(?:\.\d{1,2})?)\)\s+(.+)$")         # «8) …», «6.1) …»
REPEALED_RE = re.compile(r"^(?:\d+(?:\.\d+)*[.)]\s+)?Утратил[аио]?\s+силу", re.I)


@dataclass
class _LawUnit:
    article: str
    title: str
    url: str | None
    part: str | None
    item: str | None
    lines: list
    lead: str | None = None       # начало части перед перечнем пунктов («2. Не вправе применять …:»)

    @property
    def text(self) -> str:
        return "\n".join(self.lines).strip()


def _law_units(lines: list[str]) -> list[_LawUnit]:
    units: list[_LawUnit] = []
    article = title = url = None
    part_unit: _LawUnit | None = None
    for raw in lines:
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        if m := ARTICLE_RE.match(line):
            article, title, url, part_unit = m[1], m[2].strip(), None, None
            continue
        if article is None:
            continue
        if line.startswith("@url "):
            url = line[5:].strip()
            continue
        if m := PART_RE.match(line):
            part_unit = _LawUnit(article, title, url, m[1], None, [line])
            units.append(part_unit)
        elif (m := ITEM_RE.match(line)) and part_unit is not None:
            lead = part_unit.lines[0] if part_unit.lines[0].rstrip().endswith(":") else None
            units.append(_LawUnit(article, title, url, part_unit.part, m[1], [line], lead=lead))
        elif units and units[-1].article == article:
            units[-1].lines.append(line)            # абзац без номера — продолжение пункта
        else:
            units.append(_LawUnit(article, title, url, None, None, [line]))   # статья без частей
    return [u for u in units if u.text and not REPEALED_RE.match(u.text)]


def _range(values: list[str]) -> str | None:
    values = [v for v in dict.fromkeys(values) if v]
    if not values:
        return None
    return values[0] if len(values) == 1 else f"{values[0]}–{values[-1]}"


def _law_chunks(units: list[_LawUnit], meta: dict) -> list[Chunk]:
    chunks: list[Chunk] = []

    def emit(group: list[_LawUnit]):
        first = group[0]
        parts = [u.part for u in group]
        items = [u.item for u in group if u.item]
        text = "\n".join(u.text for u in group)
        # перечень без начала части: «Не вправе применять …: … 8) налогоплательщики…»
        if first.item and first.lead and first.lead not in text:
            text = f"{first.lead} …\n{text}"
        one_part = len(set(parts)) == 1
        # тема части статьи («1:2») или всей статьи («1»), иначе — тема документа
        keys = [f"{first.article}:{p}" for p in parts] + [first.article]
        topics = meta.get("topics") or {}
        topic = next((topics[k] for k in keys if k in topics), meta.get("topic"))
        base = dict(doc_id=meta["id"], doc_title=meta["title"], doc_kind=meta["kind"], access=meta["access"],
                    heading=first.title, page_start=None, page_end=None, article=first.article,
                    clause=_range(parts), item=_range(items) if one_part else None, short=meta.get("short"),
                    url=first.url, part_word=meta.get("part", "ч."), item_word=meta.get("item", "п."),
                    topic=topic)
        pieces = _split_long(text) if len(text) > MAX_LEN else [(0, text)]
        for k, (_, piece) in enumerate(pieces):
            if k:
                # продолжение длинной части — с её началом, иначе непонятно, о чём фрагмент
                piece = f"{text.split(chr(10))[0][:200]} …\n{piece}"
            chunks.append(Chunk(id=f"{meta['id']}#{len(chunks)}", text=piece, **base))

    buf: list[_LawUnit] = []
    for unit in units:
        size = sum(len(u.text) for u in buf)
        same = buf and buf[-1].article == unit.article
        # пункты перечня собираются со своей частью (до TARGET); разные части статьи объединяются, пока коротко
        same_part = same and buf[-1].part == unit.part and unit.item is not None
        limit = TARGET if same_part else LAW_TARGET
        if buf and (not same or size + len(unit.text) > limit or (not same_part and size > LAW_TARGET // 2)):
            emit(buf)
            buf = []
        buf.append(unit)
    if buf:
        emit(buf)
    return chunks


FAQ_Q, FAQ_SRC = "Вопрос: ", "Источник: "


def _faq_chunks(lines: list[str], meta: dict) -> list[Chunk]:
    """Вопрос ФНС — заголовок фрагмента, ответ и «Источник» (статья закона) — его текст."""
    chunks: list[Chunk] = []
    items: list[tuple[str, list[str]]] = []
    for raw in lines:
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        if line.startswith(FAQ_Q):
            items.append((line[len(FAQ_Q):].strip(), []))
        elif items:
            if line.startswith(FAQ_SRC):
                line = f"{meta.get('basis_label') or 'Основание, указанное ФНС'}: " + line[len(FAQ_SRC):].strip()
            items[-1][1].append(line)
    outdated = {re.sub(r"\s+", " ", o["question"]).strip() for o in meta.get("outdated", [])}
    for question, body in items:
        if re.sub(r"\s+", " ", question).strip() in outdated:
            continue                     # ответ устарел после изменения закона (причина — в sources.yaml)
        text = "\n".join(body)
        pieces = _split_long(text) if len(text) > MAX_LEN else [(0, text)]
        for k, (_, piece) in enumerate(pieces):
            chunks.append(Chunk(id=f"{meta['id']}#{len(chunks)}", doc_id=meta["id"], doc_title=meta["title"],
                                doc_kind=meta["kind"], access=meta["access"], clause=None,
                                heading=question if not k else f"{question} (продолжение)",
                                page_start=None, page_end=None, text=piece, short=meta.get("short"),
                                url=meta.get("url")))
    return chunks


def parse_document(path: Path, catalog: dict, stats: dict | None = None) -> list[Chunk]:
    """stats (необязательно) получает счётчики: ocr_pages — распознано страниц-сканов,
    scanned_pages — сканы, которые нечем распознать."""
    info = catalog.get(path.stem, {})
    meta = {
        "id": path.stem,
        "title": info.get("title", path.stem),
        "kind": info.get("kind", "Документ"),
        # отменённая редакция остаётся в индексе для истории, но не доступна ни одной роли
        "access": access_of(info),
        "short": info.get("short"), "url": info.get("url"), "part": info.get("part", "ч."),
        "item": info.get("item", "п."), "outdated": info.get("outdated", []), "topic": info.get("topic"),
        "topics": {str(k): v for k, v in (info.get("topics") or {}).items()},
        "basis_label": info.get("basis_label"),
    }
    if path.suffix.lower() == ".txt":
        lines = normalize(path.read_text(encoding="utf-8")).splitlines()
        return _faq_chunks(lines, meta) if info.get("structure") == "faq" else _law_chunks(_law_units(lines), meta)
    pages = _pdf_pages(path, stats) if path.suffix.lower() == ".pdf" else _docx_pages(path)
    return _chunks(_units(pages), meta)


SUFFIXES = (".pdf", ".docx", ".txt")


def parse_corpus(raw_dir: Path) -> list[Chunk]:
    catalog = load_catalog(raw_dir)
    chunks = []
    for path in sorted(raw_dir.rglob("*")):
        if path.suffix.lower() in SUFFIXES and not path.name.startswith("~$"):
            chunks.extend(parse_document(path, catalog))
    return chunks
